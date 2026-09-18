"""Genera el Excel del plan de acción con openpyxl, replicando la MISMA plantilla
visual que backend/src/modules/auditorias/export-excel.service.ts (hoja única con
franja oscura por categoría, encabezado de tabla oscuro y un bloque de resumen al
final) — incluyendo la grilla real de 48 semanas ("MESES DEL AÑO" → mes → semana
1-4) del documento origen, en vez de una celda de texto con el cronograma. Como el
plan de acción del año siguiente es una propuesta (todavía no hay ejecuciones
P/E/R/N), la celda de la semana propuesta se pinta con el color de prioridad de la
actividad en vez del semáforo de estado."""

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .periods import MESES_FULL, SEMANAS_POR_ANIO
from .report_time import formatear_generado_en
from .styles import (
    CATEGORY_FILL_HEX,
    GRID_BORDER_HEX,
    GRID_EMPTY_FILL_HEX,
    HEADER_FILL_HEX,
    MESES_BAND_FILL_HEX,
    PRIORIDAD_COLOR_HEX,
    PRIORIDAD_FILL_HEX,
    PRIORIDAD_LETRA,
)

FIXED_HEADERS = ["Actividad", "Recomendación", "Responsable", "Frecuencia (actual → propuesta)", "Prioridad"]
FIXED_WIDTHS = [30, 36, 18, 22, 10]
WEEK_COL_WIDTH = 2.7
WEEK_START = len(FIXED_HEADERS) + 1  # primera columna de la grilla de semanas
TOTAL_COLS = len(FIXED_HEADERS) + SEMANAS_POR_ANIO

_THIN_GRID = Border(*(Side(style="thin", color=GRID_BORDER_HEX) for _ in range(4)))
_HAIRLINE_BOTTOM = Border(bottom=Side(style="hair", color="E0E0E0"))


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _autosize(ws: Worksheet) -> None:
    widths = FIXED_WIDTHS + [WEEK_COL_WIDTH] * SEMANAS_POR_ANIO
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _semaforo_fill(pct: float) -> str:
    return PRIORIDAD_FILL_HEX["Bajo"] if pct >= 90 else (PRIORIDAD_FILL_HEX["Medio"] if pct >= 70 else PRIORIDAD_FILL_HEX["Alto"])


def _write_titulo(ws: Worksheet, auditoria_nombre: str, resumen: dict) -> int:
    r = 1
    titulo = ws.cell(r, 1, f"PLAN DE ACCIÓN — {auditoria_nombre.upper()} {resumen['anioPlan']}")
    titulo.font = Font(bold=True, size=14)
    titulo.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=TOTAL_COLS)
    r += 1

    generado_en = formatear_generado_en(resumen.get("generadoEn"))
    subtitulo = ws.cell(r, 1, f"Generado el {generado_en} (hora Colombia) · Análisis basado en el desempeño {resumen['anioBase']}")
    subtitulo.font = Font(italic=True, size=10, color="5B6472")
    subtitulo.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=TOTAL_COLS)
    r += 2
    return r


def _write_leyenda(ws: Worksheet, r: int) -> int:
    ws.cell(r, 1, "Leyenda de prioridad (color de la celda en la semana propuesta):").font = Font(bold=True)
    r += 1
    for i, prioridad in enumerate(("Alto", "Medio", "Bajo"), start=1):
        cell = ws.cell(r, i, f"{PRIORIDAD_LETRA[prioridad]} = {prioridad}")
        cell.font = Font(bold=True, color=PRIORIDAD_COLOR_HEX[prioridad])
        cell.fill = _fill(PRIORIDAD_FILL_HEX[prioridad])
        cell.alignment = Alignment(horizontal="center")
    return r + 2


def _write_grid_header(ws: Worksheet, r: int) -> int:
    """3 filas: encabezados fijos (combinados verticalmente) + MESES DEL AÑO / mes / semana."""
    for i, h in enumerate(FIXED_HEADERS, start=1):
        cell = ws.cell(r, i, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(HEADER_FILL_HEX)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=i, end_row=r + 2, end_column=i)

    banda = ws.cell(r, WEEK_START, "MESES DEL AÑO")
    banda.font = Font(bold=True, color="FFFFFF")
    banda.fill = _fill(MESES_BAND_FILL_HEX)
    banda.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=WEEK_START, end_row=r, end_column=WEEK_START + SEMANAS_POR_ANIO - 1)

    for m, nombre in enumerate(MESES_FULL):
        col = WEEK_START + m * 4
        cell = ws.cell(r + 1, col, nombre)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_GRID
        ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 3)

    for s in range(SEMANAS_POR_ANIO):
        col = WEEK_START + s
        cell = ws.cell(r + 2, col, (s % 4) + 1)
        cell.font = Font(size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_GRID

    return r + 3


def _write_categoria_header(ws: Worksheet, r: int, categoria: str) -> int:
    cell = ws.cell(r, 1, categoria.upper())
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = _fill(CATEGORY_FILL_HEX)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=TOTAL_COLS)
    return r + 1


def _write_actividad_row(ws: Worksheet, r: int, row: pd.Series) -> int:
    frecuencia = row["frecuencia"] if row["frecuencia"] == row["frecuencia_propuesta"] else f"{row['frecuencia']} → {row['frecuencia_propuesta']}"
    values = [row["actividad"], row["recomendacion"], row["responsable"], frecuencia]
    for i, value in enumerate(values, start=1):
        cell = ws.cell(r, i, value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = _HAIRLINE_BOTTOM

    prioridad = row["prioridad"]
    prioridad_cell = ws.cell(r, 5, prioridad)
    prioridad_cell.font = Font(bold=True, color=PRIORIDAD_COLOR_HEX[prioridad])
    prioridad_cell.fill = _fill(PRIORIDAD_FILL_HEX[prioridad])
    prioridad_cell.alignment = Alignment(horizontal="center", vertical="top")
    prioridad_cell.border = _HAIRLINE_BOTTOM

    letra = PRIORIDAD_LETRA[prioridad]
    semanas_marcadas = set(row["semanas_propuestas"])
    for s in range(SEMANAS_POR_ANIO):
        col = WEEK_START + s
        marcada = s in semanas_marcadas
        cell = ws.cell(r, col, letra if marcada else None)
        cell.border = _THIN_GRID
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if marcada:
            cell.font = Font(bold=True, size=8, color=PRIORIDAD_COLOR_HEX[prioridad])
            cell.fill = _fill(PRIORIDAD_FILL_HEX[prioridad])
        else:
            cell.fill = _fill(GRID_EMPTY_FILL_HEX)

    ws.row_dimensions[r].height = 28
    return r + 1


def _write_plan_categorias(ws: Worksheet, r: int, plan: pd.DataFrame) -> int:
    r = _write_grid_header(ws, r)  # una sola vez: se repite la banda de categoría, no la grilla de semanas
    for categoria, grupo in plan.groupby("categoria", sort=True):
        r = _write_categoria_header(ws, r, categoria)
        for _, row in grupo.iterrows():
            r = _write_actividad_row(ws, r, row)
        r += 1  # fila en blanco entre categorías
    return r


def _write_resumen(ws: Worksheet, r: int, resumen: dict, cat_metrics: pd.DataFrame) -> None:
    r += 1
    titulo = ws.cell(r, 1, f"RESUMEN / TOTALES {resumen['anioPlan']}")
    titulo.font = Font(bold=True, size=13)
    r += 2

    kpi_headers = [
        "Actividades analizadas",
        f"Cumplimiento general {resumen['anioBase']}",
        "Prioridad alta",
        "Prioridad media",
        "Prioridad baja / sin cambios",
        f"Ocurrencias propuestas {resumen['anioPlan']}",
    ]
    for i, h in enumerate(kpi_headers, start=1):
        ws.cell(r, i, h).font = Font(bold=True)
    r += 1
    kpi_values = [
        resumen["totalActividades"],
        f"{resumen['cumplimientoGeneral']}%",
        resumen["riesgoAlto"],
        resumen["riesgoMedio"],
        resumen["riesgoBajo"],
        resumen["ocurrenciasPropuestas"],
    ]
    for i, v in enumerate(kpi_values, start=1):
        ws.cell(r, i, v)
    r += 2

    ws.cell(r, 1, "Detalle por categoría").font = Font(bold=True)
    r += 1
    detalle_headers = ["Categoría", "Actividades", "Cumplimiento promedio", "Riesgo alto", "Riesgo medio"]
    for i, h in enumerate(detalle_headers, start=1):
        ws.cell(r, i, h).font = Font(bold=True)
    r += 1
    for _, row in cat_metrics.iterrows():
        ws.cell(r, 1, row["categoria"])
        ws.cell(r, 2, int(row["total_actividades"]))
        pct_cell = ws.cell(r, 3, f"{row['cumplimiento_promedio']}%")
        pct_cell.fill = _fill(_semaforo_fill(row["cumplimiento_promedio"]))
        ws.cell(r, 4, int(row["actividades_riesgo_alto"]))
        ws.cell(r, 5, int(row["actividades_riesgo_medio"]))
        r += 1


def build_excel_report(auditoria_nombre: str, resumen: dict, cat_metrics: pd.DataFrame, plan: pd.DataFrame) -> bytes:
    wb = Workbook()
    wb.creator = "Analytics Service — Plan de Acción"
    ws = wb.active
    ws.title = f"Plan de Acción {resumen['anioPlan']}"
    _autosize(ws)

    r = _write_titulo(ws, auditoria_nombre, resumen)
    r = _write_leyenda(ws, r)
    r = _write_plan_categorias(ws, r, plan)
    _write_resumen(ws, r, resumen, cat_metrics)

    ws.freeze_panes = ws.cell(1, WEEK_START).coordinate

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
