"""Genera el Excel/PDF del "Plan de Trabajo" (año en curso) que antes construía
backend/src/modules/auditorias/export-excel.service.ts y export-pdf.service.ts con
exceljs/pdfkit. Ese backend NestJS ahora solo hace de puente autenticado (ver
export-bridge.service.ts): este servicio es el único que consulta la BD (db.py) y
arma los archivos, igual que ya pasaba con el plan de acción del año siguiente
(excel_report.py / pdf_report.py)."""

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .periods import MESES_FULL, SEMANAS_POR_ANIO
from .plan_trabajo_template import (
    CATEGORY_FILL_HEX,
    ESTADO_COLOR_HEX,
    ESTADO_FILL_HEX,
    ESTADO_LETRA,
    ESTADO_NOMBRE,
    GRID_BORDER_HEX,
    GRID_EMPTY_FILL_HEX,
    HEADER_FILL_HEX,
    MESES_BAND_FILL_HEX,
    activity_week_slots,
    build_rangos_semanas,
    build_week_slots,
    cronograma_entries,
)
from .report_time import formatear_generado_en

FIXED_HEADERS = ["Actividad", "Descripción y/o evidencia", "Responsable", "Frecuencia", "% Cumpl."]
FIXED_WIDTHS = [42, 32, 22, 14, 10]
WEEK_COL_WIDTH = 2.7
WEEK_START = len(FIXED_HEADERS) + 1  # primera columna de la grilla de semanas
TOTAL_COLS = len(FIXED_HEADERS) + SEMANAS_POR_ANIO

_THIN_GRID = Border(*(Side(style="thin", color=GRID_BORDER_HEX) for _ in range(4)))
_HAIRLINE_BOTTOM = Border(bottom=Side(style="hair", color="E0E0E0"))


def _activities_by_category(df: pd.DataFrame) -> list[tuple[str, list[dict]]]:
    """DataFrame (una fila por actividad-ocurrencia, LEFT JOIN de db.fetch_activities_for_export)
    -> [(categoria, [actividad_dict, ...]), ...], preservando el orden de la consulta SQL
    (categoría, nombre de actividad, fecha programada)."""
    actividades_por_categoria: dict[str, dict[str, dict]] = {}
    orden_categorias: list[str] = []
    orden_actividades: dict[str, list[str]] = {}

    for _, row in df.iterrows():
        categoria = row["categoria"]
        if categoria not in actividades_por_categoria:
            actividades_por_categoria[categoria] = {}
            orden_categorias.append(categoria)
            orden_actividades[categoria] = []

        actividades = actividades_por_categoria[categoria]
        if row["activity_id"] not in actividades:
            actividades[row["activity_id"]] = {
                "nombre": row["actividad"],
                "descripcion": row["descripcion"] or "",
                "responsable": row["responsable"],
                "frecuencia": row["frecuencia"],
                "occurrences": [],
            }
            orden_actividades[categoria].append(row["activity_id"])

        if pd.notna(row["occurrence_id"]):
            actividades[row["activity_id"]]["occurrences"].append(
                {"fecha_programada": row["fecha_programada"], "estado": row["estado"]}
            )

    return [(categoria, [actividades_por_categoria[categoria][aid] for aid in orden_actividades[categoria]]) for categoria in orden_categorias]


def _semaforo_fill(pct: float) -> str:
    return ESTADO_FILL_HEX["EJECUTADO"] if pct >= 90 else (ESTADO_FILL_HEX["PLANEADO"] if pct >= 70 else ESTADO_FILL_HEX["NO_REALIZADO"])


# --------------------------------------------------------------------------- Excel ---


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _autosize(ws: Worksheet) -> None:
    widths = FIXED_WIDTHS + [WEEK_COL_WIDTH] * SEMANAS_POR_ANIO
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _write_titulo_excel(ws: Worksheet, auditoria_nombre: str, anio: int) -> int:
    r = 1
    titulo = ws.cell(r, 1, f"PLAN DE TRABAJO — {auditoria_nombre.upper()} {anio}")
    titulo.font = Font(bold=True, size=14)
    titulo.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=TOTAL_COLS)
    return r + 2


def _write_leyenda_excel(ws: Worksheet, r: int) -> int:
    ws.cell(r, 1, "Leyenda:").font = Font(bold=True)
    r += 1
    for i, estado in enumerate(("PLANEADO", "EJECUTADO", "REPROGRAMADO", "NO_REALIZADO"), start=1):
        cell = ws.cell(r, i, f"{ESTADO_LETRA[estado]} = {ESTADO_NOMBRE[estado]}")
        cell.font = Font(bold=True, color=ESTADO_COLOR_HEX[estado])
        cell.fill = _fill(ESTADO_FILL_HEX[estado])
        cell.alignment = Alignment(horizontal="center")
    return r + 2


def _write_grid_header(ws: Worksheet, r: int) -> int:
    """3 filas: encabezados fijos (combinados verticalmente) + MESES DEL AÑO / mes / semana."""
    for i, h in enumerate(FIXED_HEADERS, start=1):
        cell = ws.cell(r, i, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(HEADER_FILL_HEX)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=i, end_row=r + 2, end_column=i)

    banda = ws.cell(r, WEEK_START, "MESES DEL AÑO")
    banda.font = Font(bold=True, color="FFFFFF")
    banda.fill = _fill(MESES_BAND_FILL_HEX)
    banda.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=WEEK_START, end_row=r, end_column=WEEK_START + SEMANAS_POR_ANIO - 1)

    for m, nombre in enumerate(MESES_FULL):
        col = WEEK_START + m * 4
        cell = ws.cell(r + 1, col, nombre)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 3)

    for s in range(SEMANAS_POR_ANIO):
        col = WEEK_START + s
        cell = ws.cell(r + 2, col, (s % 4) + 1)
        cell.font = Font(size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    return r + 3


def _write_categoria_header_excel(ws: Worksheet, r: int, categoria: str) -> int:
    cell = ws.cell(r, 1, categoria.upper())
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = _fill(CATEGORY_FILL_HEX)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=TOTAL_COLS)
    return r + 1


def _write_actividad_row_excel(ws: Worksheet, r: int, actividad: dict) -> int:
    occurrences = actividad["occurrences"]
    total = len(occurrences)
    ejecutadas = sum(1 for o in occurrences if o["estado"] == "EJECUTADO")
    pct = f"{round(ejecutadas / total * 100)}%" if total else ""

    values = [actividad["nombre"], actividad["descripcion"], actividad["responsable"], actividad["frecuencia"], pct]
    for i, value in enumerate(values, start=1):
        cell = ws.cell(r, i, value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = _HAIRLINE_BOTTOM

    slots = activity_week_slots(occurrences)
    for s, estado in enumerate(slots):
        col = WEEK_START + s
        cell = ws.cell(r, col)
        cell.border = _THIN_GRID
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if estado:
            cell.value = ESTADO_LETRA[estado]
            cell.font = Font(bold=True, size=8, color=ESTADO_COLOR_HEX[estado])
            cell.fill = _fill(ESTADO_FILL_HEX[estado])
        else:
            cell.fill = _fill(GRID_EMPTY_FILL_HEX)

    ws.row_dimensions[r].height = 28
    return r + 1


def _write_categorias_excel(ws: Worksheet, r: int, categorias: list[tuple[str, list[dict]]]) -> tuple[int, dict, list[dict]]:
    totales = {"total": 0, "ejecutadas": 0, "reprogramadas": 0, "no_realizadas": 0}
    todas_ocurrencias: list[dict] = []
    for categoria, actividades in categorias:
        r = _write_categoria_header_excel(ws, r, categoria)
        r = _write_grid_header(ws, r)
        for actividad in actividades:
            r = _write_actividad_row_excel(ws, r, actividad)
            for o in actividad["occurrences"]:
                totales["total"] += 1
                if o["estado"] == "EJECUTADO":
                    totales["ejecutadas"] += 1
                elif o["estado"] == "REPROGRAMADO":
                    totales["reprogramadas"] += 1
                elif o["estado"] == "NO_REALIZADO":
                    totales["no_realizadas"] += 1
                todas_ocurrencias.append(o)
        r += 1  # fila en blanco entre categorías
    return r, totales, todas_ocurrencias


def _write_resumen_excel(ws: Worksheet, r: int, anio: int, totales: dict, ocurrencias: list[dict]) -> None:
    r += 1
    titulo = ws.cell(r, 1, f"RESUMEN / TOTALES {anio}")
    titulo.font = Font(bold=True, size=13)
    r += 2

    headers = ["Actividades planeadas", "Ejecutadas", "Reprogramadas", "No realizadas", "% cumplimiento general"]
    for i, h in enumerate(headers, start=1):
        ws.cell(r, i, h).font = Font(bold=True)
    r += 1
    cumplimiento_general = round(totales["ejecutadas"] / totales["total"] * 100) if totales["total"] else 0
    values = [totales["total"], totales["ejecutadas"], totales["reprogramadas"], totales["no_realizadas"], f"{cumplimiento_general}%"]
    for i, v in enumerate(values, start=1):
        ws.cell(r, i, v)
    r += 2

    ws.cell(r, 1, "Detalle por semana").font = Font(bold=True)
    r += 1
    for i, h in enumerate(["Semana", "Planeadas", "Ejecutadas", "Reprogramadas", "No realizadas"], start=1):
        ws.cell(r, i, h).font = Font(bold=True)
    r += 1
    slots = build_week_slots(ocurrencias)
    for slot in slots:
        ws.cell(r, 1, slot["label"])
        ws.cell(r, 2, slot["total"])
        ws.cell(r, 3, slot["ejecutadas"])
        ws.cell(r, 4, slot["reprogramadas"])
        ws.cell(r, 5, slot["no_realizadas"])
        r += 1
    r += 1

    ws.cell(r, 1, "% cumplimiento por rango de semanas").font = Font(bold=True)
    r += 1
    for i, h in enumerate(["Rango de semanas", "% cumplimiento"], start=1):
        ws.cell(r, i, h).font = Font(bold=True)
    r += 1
    for rango in build_rangos_semanas(slots):
        ws.cell(r, 1, rango["rango"])
        pct_cell = ws.cell(r, 2, f"{rango['cumplimiento_pct']}%")
        pct_cell.fill = _fill(_semaforo_fill(rango["cumplimiento_pct"]))
        r += 1


def build_plan_trabajo_excel(auditoria_nombre: str, anio: int, df: pd.DataFrame) -> bytes:
    wb = Workbook()
    wb.creator = "Analytics Service — Plan de Trabajo"
    ws = wb.active
    ws.title = "PLAN TRABAJO ANUAL SIG"
    _autosize(ws)

    r = _write_titulo_excel(ws, auditoria_nombre, anio)
    r = _write_leyenda_excel(ws, r)
    categorias = _activities_by_category(df)
    r, totales, ocurrencias = _write_categorias_excel(ws, r, categorias)
    _write_resumen_excel(ws, r, anio, totales, ocurrencias)

    ws.freeze_panes = ws.cell(1, WEEK_START).coordinate

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------------- PDF ---

ANCHO_TABLA = 17 * cm


def _hex(h: str) -> colors.Color:
    return colors.HexColor(f"#{h}")


def _pdf_styles() -> dict:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("PlanTitle", parent=base["Title"], fontSize=16, alignment=1, spaceAfter=4),
        "muted": ParagraphStyle("Muted", parent=base["Normal"], textColor=_hex("5B6472"), fontSize=9, alignment=1),
        "leyenda": ParagraphStyle("Leyenda", parent=base["Normal"], fontSize=9.5, textColor=_hex("555555")),
        "categoria": ParagraphStyle("Categoria", parent=base["Normal"], fontSize=10.5, textColor=colors.white, leftIndent=6),
        "actividad": ParagraphStyle("Actividad", parent=base["Normal"], fontSize=10, textColor=_hex("111111"), spaceBefore=6),
        "meta": ParagraphStyle("Meta", parent=base["Normal"], fontSize=8.5, textColor=_hex("666666")),
        "cronograma": ParagraphStyle("Cronograma", parent=base["Normal"], fontSize=8.5, textColor=_hex("333333")),
        "h2": ParagraphStyle("H2", parent=base["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=6),
    }


def _write_titulo_pdf(story: list, styles: dict, auditoria_nombre: str, anio: int, generado_en: str | None) -> None:
    story.append(Paragraph(f"PLAN DE TRABAJO — {auditoria_nombre.upper()} {anio}", styles["title"]))
    story.append(
        Paragraph(f"Generado el {formatear_generado_en(generado_en)} (hora Colombia)", styles["muted"])
    )
    story.append(Spacer(1, 10))


def _write_leyenda_pdf(story: list, styles: dict) -> None:
    partes = " &nbsp;·&nbsp; ".join(
        f'<font color="#{ESTADO_COLOR_HEX[e]}"><b>{ESTADO_LETRA[e]} = {ESTADO_NOMBRE[e]}</b></font>'
        for e in ("PLANEADO", "EJECUTADO", "REPROGRAMADO", "NO_REALIZADO")
    )
    story.append(Paragraph(f"Leyenda: {partes}", styles["leyenda"]))
    story.append(Spacer(1, 10))


def _write_categoria_header_pdf(story: list, styles: dict, categoria: str) -> None:
    banda = Table([[Paragraph(categoria.upper(), styles["categoria"])]], colWidths=[ANCHO_TABLA])
    banda.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), _hex(CATEGORY_FILL_HEX)),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(Spacer(1, 6))
    story.append(banda)
    story.append(Spacer(1, 4))


def _write_actividad_pdf(story: list, styles: dict, actividad: dict) -> None:
    occurrences = actividad["occurrences"]
    total = len(occurrences)
    ejecutadas = sum(1 for o in occurrences if o["estado"] == "EJECUTADO")
    cumplimiento = f" &nbsp;·&nbsp; {round(ejecutadas / total * 100)}% cumpl." if total else ""

    story.append(Paragraph(actividad["nombre"], styles["actividad"]))
    story.append(
        Paragraph(f"{actividad['responsable']} &nbsp;·&nbsp; {actividad['frecuencia']}{cumplimiento}", styles["meta"])
    )

    entries = cronograma_entries(occurrences)
    if entries:
        texto = "Cronograma: " + ", ".join(
            f'<font color="#{ESTADO_COLOR_HEX[e["estado"]]}"><b>{e["texto"]}</b></font>' for e in entries
        )
        story.append(Paragraph(texto, styles["cronograma"]))

    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=0.5, color=_hex("EEEEEE"), spaceBefore=0, spaceAfter=4))


def _write_actividades_pdf(story: list, styles: dict, categorias: list[tuple[str, list[dict]]]) -> tuple[dict, list[dict]]:
    totales = {"total": 0, "ejecutadas": 0, "reprogramadas": 0, "no_realizadas": 0}
    todas_ocurrencias: list[dict] = []
    for categoria, actividades in categorias:
        _write_categoria_header_pdf(story, styles, categoria)
        for actividad in actividades:
            _write_actividad_pdf(story, styles, actividad)
            for o in actividad["occurrences"]:
                totales["total"] += 1
                if o["estado"] == "EJECUTADO":
                    totales["ejecutadas"] += 1
                elif o["estado"] == "REPROGRAMADO":
                    totales["reprogramadas"] += 1
                elif o["estado"] == "NO_REALIZADO":
                    totales["no_realizadas"] += 1
                todas_ocurrencias.append(o)
    return totales, todas_ocurrencias


def _tabla_resumen(headers: list[str], data_rows: list[list], col_widths: list[float], fill_col: int | None = None) -> Table:
    data = [headers] + data_rows
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), _hex(HEADER_FILL_HEX)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E0E0E0")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if fill_col is not None:
        for i, row in enumerate(data_rows, start=1):
            pct = float(str(row[fill_col]).replace("%", ""))
            style.append(("BACKGROUND", (fill_col, i), (fill_col, i), _hex(_semaforo_fill(pct))))
    table.setStyle(TableStyle(style))
    return table


def _write_resumen_pdf(story: list, styles: dict, anio: int, totales: dict, ocurrencias: list[dict]) -> None:
    story.append(PageBreak())
    story.append(Paragraph(f"RESUMEN / TOTALES {anio}", styles["h2"]))

    cumplimiento_general = round(totales["ejecutadas"] / totales["total"] * 100) if totales["total"] else 0
    story.append(
        _tabla_resumen(
            ["Planeadas", "Ejecutadas", "Reprogramadas", "No realizadas", "% cumpl. general"],
            [[str(totales["total"]), str(totales["ejecutadas"]), str(totales["reprogramadas"]), str(totales["no_realizadas"]), f"{cumplimiento_general}%"]],
            [ANCHO_TABLA / 5] * 5,
        )
    )
    story.append(Spacer(1, 14))

    story.append(Paragraph("Detalle por semana", styles["h2"]))
    slots = build_week_slots(ocurrencias)
    story.append(
        _tabla_resumen(
            ["Semana", "Planeadas", "Ejecutadas", "Reprogramadas", "No realizadas"],
            [[s["label"], str(s["total"]), str(s["ejecutadas"]), str(s["reprogramadas"]), str(s["no_realizadas"])] for s in slots],
            [3.4 * cm, 3.4 * cm, 3.4 * cm, 3.4 * cm, 3.4 * cm],
        )
    )
    story.append(Spacer(1, 14))

    story.append(Paragraph("% cumplimiento por rango de semanas", styles["h2"]))
    rangos = build_rangos_semanas(slots)
    story.append(
        _tabla_resumen(
            ["Rango de semanas", "% cumplimiento"],
            [[r["rango"], f"{r['cumplimiento_pct']}%"] for r in rangos],
            [11 * cm, 6 * cm],
            fill_col=1,
        )
    )


def build_plan_trabajo_pdf(auditoria_nombre: str, anio: int, df: pd.DataFrame, generado_en: str | None = None) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = _pdf_styles()
    story: list = []

    _write_titulo_pdf(story, styles, auditoria_nombre, anio, generado_en)
    _write_leyenda_pdf(story, styles)

    categorias = _activities_by_category(df)
    totales, ocurrencias = _write_actividades_pdf(story, styles, categorias)
    _write_resumen_pdf(story, styles, anio, totales, ocurrencias)

    doc.build(story)
    return buffer.getvalue()
