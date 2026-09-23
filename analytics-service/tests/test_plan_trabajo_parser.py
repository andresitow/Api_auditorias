"""Pruebas de app/plan_trabajo_parser.py: lectura del Excel del plan de trabajo
(plantilla simple y plan nativo con secciones de color)."""

from datetime import date, datetime

import pytest

from app import plan_trabajo_parser as p
from app.plan_trabajo_parser import (
    CATEGORIAS_AUDITORIA,
    PlanTrabajoParseError,
    cell_text,
    is_footer_row,
    normalizar_espacios,
    normalizar_para_comparar,
    parse_activa,
    parse_frecuencia,
    parse_plan_trabajo,
    parse_plan_trabajo_bytes,
    resolver_categoria,
)
from conftest import a_bytes, crear_plan_nativo, crear_plantilla, marcar


# ------------------------------------------------------------- utilidades de texto ---


class TestCellText:
    def test_none_es_cadena_vacia(self):
        assert cell_text(None) == ""

    def test_recorta_espacios(self):
        assert cell_text("  hola  ") == "hola"

    def test_fecha_y_datetime_a_iso_solo_dia(self):
        assert cell_text(date(2026, 3, 5)) == "2026-03-05"
        assert cell_text(datetime(2026, 3, 5, 14, 30)) == "2026-03-05"

    def test_numeros_a_texto(self):
        assert cell_text(7) == "7"
        assert cell_text(1.5) == "1.5"


class TestNormalizacion:
    def test_normalizar_espacios_preserva_saltos_de_linea(self):
        assert normalizar_espacios("  a   b \n  c    d  ") == "a b\nc d"

    def test_normalizar_para_comparar_aplana_y_minuscula(self):
        assert normalizar_para_comparar("  Data   CENTER \n x ") == "data center x"

    def test_is_footer_row(self):
        assert is_footer_row("total actividades")
        assert is_footer_row("% cumplimiento")
        assert not is_footer_row("data center")
        assert not is_footer_row("")


class TestParseFrecuencia:
    @pytest.mark.parametrize(
        "raw,esperado",
        [
            ("MENSUAL", "MENSUAL"),
            ("mensual", "MENSUAL"),
            ("  Mensual  ", "MENSUAL"),
            ("DIARIA", "DIARIO"),
            ("diario", "DIARIO"),
            ("Única", "UNICA"),
            ("UNICA", "UNICA"),
            ("a demanda", "A_DEMANDA"),
            ("A_DEMANDA", "A_DEMANDA"),
            ("Cuando se requiera", "CUANDO_SE_REQUIERA"),
            ("BIMENSUAL", "BIMENSUAL"),
            ("Trimestral", "TRIMESTRAL"),
            ("semestral", "SEMESTRAL"),
            ("Anual", "ANUAL"),
        ],
    )
    def test_alias_validos(self, raw, esperado):
        assert parse_frecuencia(raw) == esperado

    @pytest.mark.parametrize("raw", ["", "quincenal", "semanal", "MENSUALES"])
    def test_invalidos_devuelven_none(self, raw):
        assert parse_frecuencia(raw) is None

    def test_todos_los_valores_coinciden_con_el_enum_del_backend(self):
        # Debe reflejar backend/prisma/schema.prisma enum Frecuencia.
        enum_backend = {
            "UNICA", "DIARIO", "MENSUAL", "BIMENSUAL", "TRIMESTRAL", "SEMESTRAL",
            "ANUAL", "A_DEMANDA", "CUANDO_SE_REQUIERA",
        }
        assert set(p.FRECUENCIA_ALIASES.values()) == enum_backend


class TestParseActiva:
    @pytest.mark.parametrize("raw", ["", "Sí", "si", "SI", "x", "1", "activa"])
    def test_verdadero(self, raw):
        assert parse_activa(raw) is True

    @pytest.mark.parametrize("raw", ["No", "no", " NO ", "false", "FALSE", "0", "inactiva", "Inactivo"])
    def test_falso(self, raw):
        assert parse_activa(raw) is False


class TestResolverCategoria:
    def test_categoria_canonica_se_mantiene(self):
        for c in CATEGORIAS_AUDITORIA:
            assert resolver_categoria(c) == c

    def test_canonica_insensible_a_mayusculas_y_espacios(self):
        assert resolver_categoria("  SERVIDORES ") == "Servidores"
        assert resolver_categoria("control  de  accesos y contraseñas") == "Control de Accesos y contraseñas"

    def test_seccion_del_plan_nativo_se_mapea(self):
        assert resolver_categoria("Data Center") == "Mantenimiento de la infraestructura"
        assert resolver_categoria("UPS") == "Mantenimiento de la infraestructura"
        assert resolver_categoria("Switch mesas usuarios") == "Switches"
        assert resolver_categoria("Implementación NIST") == "Riesgos y activos de la información"

    def test_desconocida_cae_en_otros(self):
        assert resolver_categoria("Sección nueva inventada") == "Otros"
        assert resolver_categoria("") == "Otros"

    def test_todo_valor_del_mapeo_es_categoria_canonica(self):
        for seccion, categoria in p.MAPEO_CATEGORIA_SECCION.items():
            assert categoria in CATEGORIAS_AUDITORIA, seccion

    def test_claves_del_mapeo_estan_normalizadas(self):
        for clave in p.MAPEO_CATEGORIA_SECCION:
            assert clave == normalizar_para_comparar(clave)

    def test_categorias_reflejan_las_del_backend(self):
        # backend/src/modules/auditorias/dto/create-activity.dto.ts CATEGORIAS_AUDITORIA
        assert CATEGORIAS_AUDITORIA == (
            "Sensibilización y formación SI",
            "Riesgos y activos de la información",
            "Control de Accesos y contraseñas",
            "Seguimientos como puntos de control",
            "Mantenimiento de la infraestructura",
            "Switches",
            "Servidores",
            "Otros",
        )


# ---------------------------------------------------------------- plantilla simple ---


class TestPlantillaSimple:
    def test_fila_valida(self):
        wb = crear_plantilla(
            [("Servidores", "Revisar backups", "Ana", "MENSUAL", "Captura", "obs", "Sí", None)]
        )
        res = parse_plan_trabajo(wb)
        assert res.formato == "plantilla"
        assert res.anio == datetime.now().year
        assert res.errores == []
        assert len(res.filas) == 1
        f = res.filas[0]
        assert f.fila == 2
        assert f.categoria == "Servidores"
        assert f.nombre == "Revisar backups"
        assert f.responsable == "Ana"
        assert f.frecuencia == "MENSUAL"
        assert f.descripcionEvidencia == "Captura"
        assert f.observacion == "obs"
        assert f.activa is True
        assert f.fechaEspecifica is None
        assert f.periodos == []  # la plantilla simple no trae estados

    def test_encabezado_sin_tilde_tambien_se_reconoce(self):
        wb = crear_plantilla([("Servidores", "X", "Ana", "MENSUAL", None, None, None, None)], encabezado="categoria")
        assert parse_plan_trabajo(wb).formato == "plantilla"

    def test_categoria_se_lleva_a_una_de_las_8(self):
        wb = crear_plantilla([("Data Center", "X", "Ana", "MENSUAL", None, None, None, None)])
        assert parse_plan_trabajo(wb).filas[0].categoria == "Mantenimiento de la infraestructura"

    def test_categoria_desconocida_cae_en_otros(self):
        wb = crear_plantilla([("Rara", "X", "Ana", "MENSUAL", None, None, None, None)])
        assert parse_plan_trabajo(wb).filas[0].categoria == "Otros"

    def test_filas_vacias_se_ignoran(self):
        wb = crear_plantilla(
            [
                ("Servidores", "A", "Ana", "MENSUAL", None, None, None, None),
                (None, None, None, None, None, None, None, None),
                ("Switches", "B", "Luis", "ANUAL", None, None, None, None),
            ]
        )
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["A", "B"]
        assert res.errores == []

    def test_falta_categoria(self):
        wb = crear_plantilla([(None, "A", "Ana", "MENSUAL", None, None, None, None)])
        res = parse_plan_trabajo(wb)
        assert res.filas == []
        assert [(e.fila, e.motivo) for e in res.errores] == [(2, "Falta la categoría")]

    def test_falta_nombre_responsable_o_frecuencia_invalida(self):
        wb = crear_plantilla(
            [
                ("Servidores", None, "Ana", "MENSUAL", None, None, None, None),
                ("Servidores", "A", None, "MENSUAL", None, None, None, None),
                ("Servidores", "B", "Ana", "quincenal", None, None, None, None),
            ]
        )
        res = parse_plan_trabajo(wb)
        assert res.filas == []
        assert [(e.fila, e.motivo) for e in res.errores] == [
            (2, "Falta el nombre de la actividad"),
            (3, "Falta el responsable"),
            (4, 'Frecuencia inválida: "quincenal"'),
        ]

    def test_unica_requiere_fecha_especifica(self):
        wb = crear_plantilla(
            [
                ("Servidores", "A", "Ana", "UNICA", None, None, None, None),
                ("Servidores", "B", "Ana", "Única", None, None, None, "2026-09-15"),
            ]
        )
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["B"]
        assert res.filas[0].fechaEspecifica == "2026-09-15"
        assert res.errores[0].motivo == "La frecuencia UNICA requiere fecha específica"

    def test_fecha_como_datetime_de_excel_se_normaliza_a_iso(self):
        wb = crear_plantilla(
            [("Servidores", "A", "Ana", "UNICA", None, None, None, datetime(2026, 9, 15, 0, 0))]
        )
        assert parse_plan_trabajo(wb).filas[0].fechaEspecifica == "2026-09-15"

    def test_columna_activa(self):
        wb = crear_plantilla(
            [
                ("Servidores", "A", "Ana", "MENSUAL", None, None, "No", None),
                ("Servidores", "B", "Ana", "MENSUAL", None, None, "Sí", None),
                ("Servidores", "C", "Ana", "MENSUAL", None, None, None, None),
            ]
        )
        assert [f.activa for f in parse_plan_trabajo(wb).filas] == [False, True, True]

    def test_descripcion_y_observacion_vacias_son_none(self):
        wb = crear_plantilla([("Servidores", "A", "Ana", "MENSUAL", "", "", None, None)])
        f = parse_plan_trabajo(wb).filas[0]
        assert f.descripcionEvidencia is None
        assert f.observacion is None

    def test_un_error_no_aborta_las_demas_filas(self):
        wb = crear_plantilla(
            [
                ("Servidores", "A", "Ana", "MENSUAL", None, None, None, None),
                ("Servidores", "B", None, "MENSUAL", None, None, None, None),
                ("Servidores", "C", "Ana", "MENSUAL", None, None, None, None),
            ]
        )
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["A", "C"]
        assert len(res.errores) == 1


# ------------------------------------------------------------------- plan nativo ---


class TestPlanNativoEstructura:
    def test_detecta_formato_nativo_y_anio_del_titulo(self, act):
        wb = crear_plan_nativo([("Data Center", [act("Revisión")])], anio=2027)
        res = parse_plan_trabajo(wb)
        assert res.formato == "nativo"
        assert res.anio == 2027

    def test_sin_anio_en_titulo_usa_el_anio_actual(self, act):
        wb = crear_plan_nativo([("Data Center", [act("Revisión")])], titulo="PLAN SIN AÑO")
        assert parse_plan_trabajo(wb).anio == datetime.now().year

    def test_las_actividades_heredan_la_categoria_de_su_seccion(self, act):
        wb = crear_plan_nativo(
            [
                ("Data Center", [act("A"), act("B")]),
                ("Switch mesas usuarios", [act("C")]),
                ("Sección inventada", [act("D")]),
            ]
        )
        res = parse_plan_trabajo(wb)
        assert [(f.nombre, f.categoria) for f in res.filas] == [
            ("A", "Mantenimiento de la infraestructura"),
            ("B", "Mantenimiento de la infraestructura"),
            ("C", "Switches"),
            ("D", "Otros"),
        ]

    def test_lee_columnas_a_d(self, act):
        wb = crear_plan_nativo(
            [("UPS", [act("Mantenimiento UPS", responsable="Paula", frecuencia="Semestral", descripcion="Acta")])]
        )
        f = parse_plan_trabajo(wb).filas[0]
        assert (f.nombre, f.descripcionEvidencia, f.responsable, f.frecuencia) == (
            "Mantenimiento UPS",
            "Acta",
            "Paula",
            "SEMESTRAL",
        )
        assert f.activa is True
        assert f.fechaEspecifica is None

    def test_numero_de_fila_es_el_de_excel(self, act):
        wb = crear_plan_nativo([("UPS", [act("A"), act("B")])])
        # fila 1-2 título, fila 3 sección, actividades en 4 y 5
        assert [f.fila for f in parse_plan_trabajo(wb).filas] == [4, 5]

    def test_las_filas_de_encabezado_y_totales_no_son_actividades(self, act):
        wb = crear_plan_nativo([("UPS", [act("A")])], footer=True)
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["A"]
        assert res.errores == []

    def test_fila_de_total_suelta_sin_fusion_se_ignora(self, act):
        wb = crear_plan_nativo([("UPS", [act("A")])], footer=False)
        ws = wb.active
        ws.cell(5, 1, "Total verificación")  # sin fusión A:D
        ws.cell(5, 3, "algo")
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["A"]
        assert res.errores == []

    def test_filas_vacias_entre_actividades_se_ignoran(self, act):
        wb = crear_plan_nativo([("UPS", [act("A"), act(None, None, None), act("B")])])
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["A", "B"]
        assert res.errores == []

    def test_filas_antes_de_la_primera_categoria_se_ignoran(self, act):
        wb = crear_plan_nativo(
            [("UPS", [act("A")])],
            filas_previas=[("texto suelto", None, "Alguien", "MENSUAL")],
        )
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["A"]
        assert res.errores == []

    def test_errores_de_validacion_por_fila(self, act):
        wb = crear_plan_nativo(
            [
                (
                    "UPS",
                    [
                        act("A", responsable=None),
                        act("B", frecuencia="quincenal"),
                        act(None, responsable="Ana"),
                        act("OK"),
                    ],
                )
            ]
        )
        res = parse_plan_trabajo(wb)
        assert [f.nombre for f in res.filas] == ["OK"]
        assert [e.motivo for e in res.errores] == [
            "Falta el responsable",
            'Frecuencia inválida: "quincenal"',
            "Falta el nombre de la actividad",
        ]

    def test_sin_ninguna_categoria_lanza_error_de_formato(self):
        wb = crear_plan_nativo([])
        with pytest.raises(PlanTrabajoParseError):
            parse_plan_trabajo(wb)

    def test_solo_una_hoja_vacia_lanza_error(self):
        from openpyxl import Workbook

        with pytest.raises(PlanTrabajoParseError):
            parse_plan_trabajo(Workbook())

    def test_solo_se_lee_la_primera_hoja(self, act):
        wb = crear_plan_nativo([("UPS", [act("A")])])
        otra = wb.create_sheet("Otra")
        otra["A1"] = "categoría"
        assert parse_plan_trabajo(wb).formato == "nativo"

    # CARACTERIZACIÓN: en el plan nativo no existe columna de fecha, así que una actividad
    # UNICA siempre se reporta como error. Consecuencia en el backend: al reimportar, la
    # actividad UNICA existente no queda en `vistas` y se remueve como "sobrante".
    def test_unica_en_plan_nativo_siempre_es_error(self, act):
        wb = crear_plan_nativo([("UPS", [act("Auditoría externa", frecuencia="UNICA")])])
        res = parse_plan_trabajo(wb)
        assert res.filas == []
        assert res.errores[0].motivo == "La frecuencia UNICA requiere fecha específica"


class TestPlanNativoPeriodos:
    """Traducción de la grilla H:BC (12 meses x 4 semanas) a periodos del backend."""

    def _periodos(self, act, frecuencia, marcas):
        wb = crear_plan_nativo([("UPS", [act("A", frecuencia=frecuencia, marcas=marcas)])])
        return parse_plan_trabajo(wb).filas[0].periodos

    def test_mensual_siempre_devuelve_12_periodos_con_p_por_defecto(self, act):
        per = self._periodos(act, "MENSUAL", {})
        assert [x.periodo for x in per] == [f"2026-{m:02d}" for m in range(1, 13)]
        assert {x.estado for x in per} == {"PLANEADO"}
        assert not any(x.esAdHoc for x in per)

    def test_mensual_respeta_el_codigo_de_cada_mes(self, act):
        per = self._periodos(act, "MENSUAL", {(0, 0): "E", (2, 1): "R", (4, 3): "N", (5, 2): "P"})
        estados = {x.periodo: x.estado for x in per}
        assert estados["2026-01"] == "EJECUTADO"
        assert estados["2026-02"] == "PLANEADO"
        assert estados["2026-03"] == "REPROGRAMADO"
        assert estados["2026-05"] == "NO_REALIZADO"
        assert estados["2026-06"] == "PLANEADO"

    def test_codigos_en_minuscula_y_con_espacios_se_aceptan(self, act):
        per = self._periodos(act, "MENSUAL", {(0, 0): " e "})
        assert per[0].estado == "EJECUTADO"

    def test_valores_no_reconocidos_se_ignoran(self, act):
        per = self._periodos(act, "MENSUAL", {(0, 0): "X", (1, 0): "OK", (2, 0): "1"})
        assert {x.estado for x in per} == {"PLANEADO"}

    def test_dentro_de_un_mes_gana_el_primer_codigo_reconocible(self, act):
        per = self._periodos(act, "MENSUAL", {(0, 1): "N", (0, 3): "E"})
        assert per[0].estado == "NO_REALIZADO"

    def test_bimensual_6_periodos(self, act):
        per = self._periodos(act, "BIMENSUAL", {(0, 0): "E"})
        assert [x.periodo for x in per] == [f"2026-B{b}" for b in range(1, 7)]
        assert per[0].estado == "EJECUTADO"
        assert per[1].estado == "PLANEADO"

    def test_trimestral_4_periodos(self, act):
        per = self._periodos(act, "TRIMESTRAL", {})
        assert [x.periodo for x in per] == ["2026-Q1", "2026-Q2", "2026-Q3", "2026-Q4"]

    def test_semestral_2_periodos(self, act):
        per = self._periodos(act, "SEMESTRAL", {(7, 0): "E"})
        assert [(x.periodo, x.estado) for x in per] == [("2026-S1", "PLANEADO"), ("2026-S2", "EJECUTADO")]

    def test_anual_un_periodo(self, act):
        per = self._periodos(act, "ANUAL", {(11, 3): "R"})
        assert [(x.periodo, x.estado) for x in per] == [("2026", "REPROGRAMADO")]

    def test_prioridad_ejecutado_sobre_reprogramado_sobre_no_realizado_sobre_planeado(self, act):
        # Q1: P + E + N -> E ; Q2: N + R -> R ; Q3: P + N -> N ; Q4: solo P -> P
        marcas = {
            (0, 0): "P", (1, 0): "E", (2, 0): "N",
            (3, 0): "N", (4, 0): "R",
            (6, 0): "P", (7, 0): "N",
            (9, 0): "P",
        }
        per = self._periodos(act, "TRIMESTRAL", marcas)
        assert [x.estado for x in per] == ["EJECUTADO", "REPROGRAMADO", "NO_REALIZADO", "PLANEADO"]

    def test_prioridad_no_depende_del_orden_cronologico(self, act):
        # La E está en el primer mes y una P en el último: gana E (no "la última celda").
        per = self._periodos(act, "TRIMESTRAL", {(0, 0): "E", (2, 0): "P"})
        assert per[0].estado == "EJECUTADO"

    def test_a_demanda_solo_emite_meses_con_marca_y_esAdHoc(self, act):
        per = self._periodos(act, "A_DEMANDA", {(1, 0): "E", (6, 2): "N"})
        assert [(x.periodo, x.estado, x.esAdHoc) for x in per] == [
            ("2026-02", "EJECUTADO", True),
            ("2026-07", "NO_REALIZADO", True),
        ]

    def test_cuando_se_requiera_igual_que_a_demanda(self, act):
        per = self._periodos(act, "Cuando se requiera", {(3, 0): "P"})
        assert [(x.periodo, x.estado, x.esAdHoc) for x in per] == [("2026-04", "PLANEADO", True)]

    def test_a_demanda_sin_marcas_no_emite_periodos(self, act):
        assert self._periodos(act, "A_DEMANDA", {}) == []

    def test_usa_el_anio_del_titulo_en_los_identificadores(self, act):
        wb = crear_plan_nativo([("UPS", [act("A", frecuencia="ANUAL")])], anio=2031)
        assert parse_plan_trabajo(wb).filas[0].periodos[0].periodo == "2031"

    def test_los_identificadores_coinciden_con_periods_util_del_backend(self, act):
        # backend/src/modules/auditorias/periods.util.ts genera exactamente estos ids.
        from app.periods import periods_for_year

        for frecuencia in ("MENSUAL", "BIMENSUAL", "TRIMESTRAL", "SEMESTRAL", "ANUAL"):
            per = self._periodos(act, frecuencia, {})
            assert [x.periodo for x in per] == [pid for pid, _ in periods_for_year(frecuencia, 2026)]


class TestPlanNativoDiario:
    """DIARIO se lee semana a semana (48 columnas) y se mapea a periodos AAAA-Wnn."""

    def _periodos(self, act, marcas, anio=2026):
        wb = crear_plan_nativo([("UPS", [act("Diario", frecuencia="DIARIO", marcas=marcas)])], anio=anio)
        return parse_plan_trabajo(wb).filas[0].periodos

    def test_sin_marcas_no_devuelve_periodos(self, act):
        assert self._periodos(act, {}) == []

    @pytest.mark.parametrize(
        "mes,semana,esperado",
        [
            # enero (31 días): punto medio de la franja -> día aprox 4, 12, 19, 27
            (0, 0, "2026-W01"),
            (0, 1, "2026-W02"),
            (0, 2, "2026-W03"),
            (0, 3, "2026-W04"),
        ],
    )
    def test_enero_mapea_a_las_semanas_1_a_4(self, act, mes, semana, esperado):
        per = self._periodos(act, {(mes, semana): "E"})
        assert [(x.periodo, x.estado) for x in per] == [(esperado, "EJECUTADO")]

    def test_cada_marca_genera_su_propio_periodo(self, act):
        per = self._periodos(act, {(0, 0): "E", (0, 1): "E", (0, 2): "N", (0, 3): "P"})
        assert [(x.periodo, x.estado) for x in per] == [
            ("2026-W01", "EJECUTADO"),
            ("2026-W02", "EJECUTADO"),
            ("2026-W03", "NO_REALIZADO"),
            ("2026-W04", "PLANEADO"),
        ]

    def test_no_marca_esAdHoc(self, act):
        assert not any(x.esAdHoc for x in self._periodos(act, {(0, 0): "E"}))

    def test_diciembre_cae_dentro_de_las_semanas_del_sistema(self, act):
        per = self._periodos(act, {(11, 3): "E"})
        semana = int(per[0].periodo.split("-W")[1])
        assert 49 <= semana <= 53

    def test_todos_los_periodos_generados_existen_en_periods_util(self, act):
        from app.periods import periods_for_year

        validos = {pid for pid, _ in periods_for_year("DIARIO", 2026)}
        marcas = {(m, s): "E" for m in range(12) for s in range(4)}
        per = self._periodos(act, marcas)
        assert len(per) == 48
        assert {x.periodo for x in per} <= validos

    def test_es_anio_bisiesto_seguro(self, act):
        per = self._periodos(act, {(1, 3): "E"}, anio=2024)  # febrero 2024 tiene 29 días
        assert per[0].periodo.startswith("2024-W")


# ------------------------------------------------------------------- bytes / API ---


class TestParseBytes:
    def test_desde_bytes_de_un_xlsx(self, act):
        contenido = a_bytes(crear_plan_nativo([("UPS", [act("A")])]))
        res = parse_plan_trabajo_bytes(contenido)
        assert res.formato == "nativo"
        assert len(res.filas) == 1

    def test_bytes_que_no_son_xlsx_lanzan_excepcion(self):
        with pytest.raises(Exception):
            parse_plan_trabajo_bytes(b"esto no es un excel")


def test_marcar_helper_escribe_en_la_columna_correcta():
    # Guardia de la propia utilería de pruebas: mes 0/semana 0 = columna H (8).
    from openpyxl import Workbook

    ws = Workbook().active
    marcar(ws, 5, 0, 0, "E")
    marcar(ws, 5, 11, 3, "N")
    assert ws.cell(5, 8).value == "E"
    assert ws.cell(5, 55).value == "N"  # BC = columna 55
