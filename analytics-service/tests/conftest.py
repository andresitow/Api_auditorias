"""Configuración común de las pruebas del analytics-service.

Las variables de entorno se fijan ANTES de importar `app.*` porque `app/config.py`
las exige al importarse (y `app/db.py` crea el engine de SQLAlchemy en el import;
create_engine es perezoso, así que nunca se abre una conexión real a Postgres)."""

import io
import os

os.environ["DATABASE_URL"] = "postgresql://test:test@localhost:5432/test_no_real_db"
os.environ["JWT_SECRET"] = "secreto-solo-para-pruebas"
os.environ.setdefault("CORS_ORIGINS", "http://localhost:3000")

import pytest  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.utils import column_index_from_string  # noqa: E402

WEEK_FIRST = column_index_from_string("H")  # primera columna de la grilla H:BC


def marcar(ws, fila: int, mes_index0: int, semana_index0: int, codigo: str) -> None:
    """Escribe un código en la grilla mensual: 12 meses x 4 semanas desde la columna H."""
    ws.cell(fila, WEEK_FIRST + mes_index0 * 4 + semana_index0, codigo)


def crear_plan_nativo(
    secciones, anio: int = 2026, titulo: str | None = None, footer: bool = True, filas_previas=()
) -> Workbook:
    """Arma un libro con el formato "PLAN TRABAJO ANUAL SIG".

    `secciones` = [(titulo_seccion, [actividad, ...]), ...] donde cada actividad es un
    dict con nombre, descripcion, responsable, frecuencia y `marcas` = {(mes0, semana0): "E"}.
    La fila del título (B1:BC2 fusionada) lleva el año; cada sección es una fusión A:D de
    una sola fila (igual que en el archivo real)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "PLAN TRABAJO ANUAL SIG"
    ws["A1"] = "SIG"
    ws["B1"] = titulo if titulo is not None else f"PLAN DE TRABAJO ANUAL {anio}"
    ws.merge_cells("B1:BC2")

    fila = 3
    for valores in filas_previas:  # filas sueltas antes de la primera categoría
        for col, valor in enumerate(valores, start=1):
            ws.cell(fila, col, valor)
        fila += 1
    for titulo_seccion, actividades in secciones:
        ws.cell(fila, 1, titulo_seccion)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        fila += 1
        for act in actividades:
            ws.cell(fila, 1, act.get("nombre"))
            ws.cell(fila, 2, act.get("descripcion"))
            ws.cell(fila, 3, act.get("responsable"))
            ws.cell(fila, 4, act.get("frecuencia"))
            for (mes, semana), codigo in act.get("marcas", {}).items():
                marcar(ws, fila, mes, semana, codigo)
            fila += 1
    if footer:
        ws.cell(fila, 1, "TOTAL ACTIVIDADES")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        ws.cell(fila + 1, 1, "% cumplimiento")
        ws.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 1, end_column=4)
    return wb


def crear_plantilla(filas, encabezado: str = "Categoría") -> Workbook:
    """Plantilla simple: A=Categoría B=Actividad C=Responsable D=Frecuencia E=Descripción
    F=Observación G=Activa H=Fecha específica."""
    wb = Workbook()
    ws = wb.active
    ws.append([encabezado, "Actividad", "Responsable", "Frecuencia", "Descripción", "Observación", "Activa", "Fecha"])
    for f in filas:
        ws.append(list(f))
    return wb


def a_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def act():
    """Fábrica de actividades para crear_plan_nativo."""

    def _act(nombre="Actividad", responsable="Ana", frecuencia="MENSUAL", descripcion=None, marcas=None):
        return {
            "nombre": nombre,
            "descripcion": descripcion,
            "responsable": responsable,
            "frecuencia": frecuencia,
            "marcas": marcas or {},
        }

    return _act
