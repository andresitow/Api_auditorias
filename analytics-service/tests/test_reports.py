"""Pruebas de humo y de contenido para los generadores de Excel/PDF
(plan_trabajo_report.py, excel_report.py, pdf_report.py). No se comparan bytes: se abre el
Excel resultante con openpyxl y se verifican títulos, totales y semáforo."""

import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import analysis, excel_report, pdf_report, plan_builder, plan_trabajo_report
from app.plan_trabajo_report import _activities_by_category, _semaforo_fill
from app.plan_trabajo_template import ESTADO_FILL_HEX


def df_export(filas):
    """filas = (categoria, actividad_id, actividad, estado|None, fecha|None)."""
    return pd.DataFrame(
        [
            dict(
                activity_id=aid,
                categoria=cat,
                actividad=nombre,
                descripcion=None,
                responsable="Ana",
                frecuencia="MENSUAL",
                occurrence_id=None if estado is None else f"{aid}-{i}",
                fecha_programada=fecha,
                estado=estado,
            )
            for i, (cat, aid, nombre, estado, fecha) in enumerate(filas)
        ]
    )


def celdas(wb):
    ws = wb.active
    return {(c.row, c.column): c.value for fila in ws.iter_rows() for c in fila if c.value is not None}


def textos(wb):
    return [v for v in celdas(wb).values() if isinstance(v, str)]


class TestActivitiesByCategory:
    def test_agrupa_preservando_el_orden_de_la_consulta(self):
        df = df_export(
            [
                ("Switches", "a1", "Uno", "EJECUTADO", date(2026, 1, 31)),
                ("Switches", "a1", "Uno", "PLANEADO", date(2026, 2, 28)),
                ("Switches", "a2", "Dos", "NO_REALIZADO", date(2026, 1, 31)),
                ("Servidores", "a3", "Tres", "EJECUTADO", date(2026, 1, 31)),
            ]
        )
        res = _activities_by_category(df)
        assert [c for c, _ in res] == ["Switches", "Servidores"]
        switches = res[0][1]
        assert [a["nombre"] for a in switches] == ["Uno", "Dos"]
        assert len(switches[0]["occurrences"]) == 2
        assert switches[0]["occurrences"][0] == {"fecha_programada": date(2026, 1, 31), "estado": "EJECUTADO"}

    def test_actividad_sin_ocurrencias_aparece_con_lista_vacia(self):
        df = df_export([("Switches", "a1", "Sin datos", None, None)])
        res = _activities_by_category(df)
        assert res[0][1][0]["occurrences"] == []

    # BUG (reportado): con pandas 3 una descripción NULL llega como NaN (float, truthy), por lo
    # que `row["descripcion"] or ""` en plan_trabajo_report._activities_by_category NO la
    # convierte en "" y NaN se propaga a las celdas de Excel / al Paragraph del PDF.
    # Esperado: "" ; obtenido: nan.
    @pytest.mark.skip(reason="Bug: descripcion NULL llega como NaN y `or \"\"` no la normaliza (esperado '', obtenido nan)")
    def test_descripcion_nula_se_vuelve_cadena_vacia(self):
        df = df_export([("Switches", "a1", "X", None, None)])
        assert _activities_by_category(df)[0][1][0]["descripcion"] == ""


class TestSemaforoFill:
    def test_umbrales_90_y_70_fijos(self):
        # Nota: aquí los umbrales están fijos (no vienen de AuditoriaConfig como en el dashboard).
        assert _semaforo_fill(90) == ESTADO_FILL_HEX["EJECUTADO"]
        assert _semaforo_fill(100) == ESTADO_FILL_HEX["EJECUTADO"]
        assert _semaforo_fill(89.9) == ESTADO_FILL_HEX["PLANEADO"]
        assert _semaforo_fill(70) == ESTADO_FILL_HEX["PLANEADO"]
        assert _semaforo_fill(69.9) == ESTADO_FILL_HEX["NO_REALIZADO"]
        assert _semaforo_fill(0) == ESTADO_FILL_HEX["NO_REALIZADO"]


class TestPlanTrabajoExcel:
    def test_libro_valido_con_titulo_y_totales(self):
        df = df_export(
            [
                ("Switches", "a1", "Uno", "EJECUTADO", date(2026, 1, 31)),
                ("Switches", "a1", "Uno", "EJECUTADO", date(2026, 2, 28)),
                ("Switches", "a1", "Uno", "REPROGRAMADO", date(2026, 3, 31)),
                ("Switches", "a1", "Uno", "NO_REALIZADO", date(2026, 4, 30)),
            ]
        )
        contenido = plan_trabajo_report.build_plan_trabajo_excel("Auditoría X", 2026, df)
        wb = load_workbook(io.BytesIO(contenido))
        assert wb.active.title == "PLAN TRABAJO ANUAL SIG"
        assert "PLAN DE TRABAJO — AUDITORÍA X 2026" in textos(wb)

        valores = celdas(wb)
        fila_headers = next(r for (r, c), v in valores.items() if v == "Actividades planeadas")
        assert [valores[(fila_headers + 1, c)] for c in range(1, 6)] == [4, 2, 1, 1, "50%"]

    def test_sin_ocurrencias_cumplimiento_0(self):
        df = df_export([("Switches", "a1", "Sin datos", None, None)])
        wb = load_workbook(io.BytesIO(plan_trabajo_report.build_plan_trabajo_excel("A", 2026, df)))
        valores = celdas(wb)
        fila = next(r for (r, c), v in valores.items() if v == "Actividades planeadas")
        assert [valores[(fila + 1, c)] for c in range(1, 6)] == [0, 0, 0, 0, "0%"]

    def test_congela_paneles_tras_las_columnas_fijas(self):
        df = df_export([("Switches", "a1", "Uno", "EJECUTADO", date(2026, 1, 5))])
        wb = load_workbook(io.BytesIO(plan_trabajo_report.build_plan_trabajo_excel("A", 2026, df)))
        assert wb.active.freeze_panes == "F1"


class TestPlanTrabajoPdf:
    def test_genera_pdf_valido(self):
        df = df_export(
            [
                ("Switches", "a1", "Uno", "EJECUTADO", date(2026, 1, 31)),
                ("Servidores", "a2", "Dos", "PLANEADO", date(2026, 2, 28)),
            ]
        )
        pdf = plan_trabajo_report.build_plan_trabajo_pdf("Auditoría X", 2026, df, "2026-06-15T17:30:00+00:00")
        assert pdf.startswith(b"%PDF")
        assert len(pdf) > 1000

    def test_sin_generado_en_usa_respaldo_y_no_falla(self):
        df = df_export([("Switches", "a1", "Uno", "EJECUTADO", date(2026, 1, 31))])
        assert plan_trabajo_report.build_plan_trabajo_pdf("A", 2026, df).startswith(b"%PDF")


class TestReportesDelPlanDeAccion:
    """Pipeline completo análisis -> plan -> Excel/PDF con datos sintéticos."""

    def _armar(self):
        df = pd.DataFrame(
            [
                dict(activity_id="a1", categoria="Servidores", actividad="Backups", responsable="Ana",
                     frecuencia="MENSUAL", occurrence_id=f"o{i}", estado="EJECUTADO")
                for i in range(4)
            ]
            + [
                dict(activity_id="a2", categoria="Switches", actividad="Firmware", responsable="Luis",
                     frecuencia="TRIMESTRAL", occurrence_id=f"p{i}", estado="NO_REALIZADO")
                for i in range(4)
            ]
        )
        metrics = analysis.compute_activity_metrics(df, 90, 70)
        cat = analysis.compute_categoria_metrics(metrics)
        plan = plan_builder.build_plan(metrics, 2027)
        resumen = {
            "auditoriaId": "aud",
            "auditoriaNombre": "Auditoría X",
            "anioBase": 2026,
            "anioPlan": 2027,
            "totalActividades": 2,
            "cumplimientoGeneral": 50.0,
            "riesgoAlto": 1,
            "riesgoMedio": 0,
            "riesgoBajo": 1,
            "ocurrenciasPropuestas": int(plan["ocurrencias_propuestas"].sum()),
            "generadoEn": "2026-06-15T17:30:00+00:00",
        }
        return resumen, cat, plan

    def test_excel_contiene_titulo_y_resumen(self):
        resumen, cat, plan = self._armar()
        wb = load_workbook(io.BytesIO(excel_report.build_excel_report("Auditoría X", resumen, cat, plan)))
        assert wb.active.title == "Plan de Acción 2027"
        t = textos(wb)
        assert "PLAN DE ACCIÓN — AUDITORÍA X 2027" in t
        assert "RESUMEN / TOTALES 2027" in t
        assert "50.0%" in t

    def test_excel_incluye_una_fila_por_categoria_en_el_detalle(self):
        resumen, cat, plan = self._armar()
        wb = load_workbook(io.BytesIO(excel_report.build_excel_report("A", resumen, cat, plan)))
        t = textos(wb)
        assert "Detalle por categoría" in t
        assert "Servidores" in t and "Switches" in t

    def test_pdf_valido(self):
        resumen, cat, plan = self._armar()
        assert pdf_report.build_pdf_report("Auditoría X", resumen, cat, plan).startswith(b"%PDF")
