"""Pruebas de scripts/plan_trabajo_sync.py (limpieza/sincronización in-place del Excel real)."""

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

import plan_trabajo_sync as s  # noqa: E402
from conftest import crear_plan_nativo, marcar  # noqa: E402


class TestNormalizarTexto:
    def test_no_str_no_se_toca(self):
        assert s.normalizar_texto(None) is None
        assert s.normalizar_texto(5) is None

    def test_texto_ya_limpio_devuelve_none(self):
        assert s.normalizar_texto("Hola mundo") is None

    def test_colapsa_espacios(self):
        assert s.normalizar_texto("Hola    mundo  ") == "Hola mundo"

    def test_preserva_saltos_internos_y_recorta_los_de_los_bordes(self):
        assert s.normalizar_texto("\n a   b \n c \n") == "a b\nc"


class TestNormalizarEstado:
    def test_none(self):
        assert s.normalizar_estado(None) == (None, None)

    def test_minuscula_se_normaliza(self):
        assert s.normalizar_estado("e") == ("E", None)
        assert s.normalizar_estado(" n ") == ("N", None)

    def test_ya_correcto_no_cambia(self):
        assert s.normalizar_estado("P") == (None, None)

    def test_invalido_da_advertencia_y_no_adivina(self):
        valor, adv = s.normalizar_estado("X")
        assert valor is None and "no es un código P/E/R/N válido" in adv

    def test_no_texto_da_advertencia(self):
        valor, adv = s.normalizar_estado(1)
        assert valor is None and "no-texto" in adv

    def test_solo_espacios_se_vacia(self):
        assert s.normalizar_estado("  ") == ("", None) or s.normalizar_estado("  ") == (None, None)


class TestDetectarEstructura:
    def test_detecta_categorias_totales_y_rango(self, act):
        wb = crear_plan_nativo([("UPS", [act("A"), act("B")]), ("Data Center", [act("C")])])
        ws = wb.active
        est = s.detectar_estructura(ws)
        assert est.categoria_rows == {3, 6}
        assert est.filas_totales == {8, 9}
        assert est.primera_fila == 3
        assert est.ultima_fila == 7  # justo antes de los totales

    def test_sin_totales_llega_hasta_el_final(self, act):
        wb = crear_plan_nativo([("UPS", [act("A")])], footer=False)
        ws = wb.active
        est = s.detectar_estructura(ws)
        assert est.filas_totales == set()
        assert est.ultima_fila == ws.max_row

    def test_sin_categorias_lanza_valueerror(self):
        with pytest.raises(ValueError):
            s.detectar_estructura(Workbook().active)


class TestSincronizar:
    def _ws(self, act, marcas=None, nombre="Actividad", descr="desc"):
        wb = crear_plan_nativo(
            [("UPS", [act(nombre, frecuencia="MENSUAL", descripcion=descr, marcas=marcas or {})])]
        )
        return wb.active

    def test_limpia_espacios_y_codigos_sin_tocar_lo_correcto(self, act):
        ws = self._ws(act, marcas={(0, 0): "e", (1, 0): "P"}, nombre="  Revisión   general ")
        est = s.detectar_estructura(ws)
        cambios, adv = s.sincronizar(ws, est)
        celdas = {c.celda: (c.antes, c.despues) for c in cambios}
        assert celdas["A4"] == ("  Revisión   general ", "Revisión general")
        assert celdas["H4"] == ("e", "E")
        assert "L4" not in celdas  # "P" ya correcto
        assert adv == []
        assert ws["H4"].value == "E"

    def test_no_toca_filas_de_categoria_ni_de_totales(self, act):
        ws = self._ws(act)
        ws["A3"].value = "  UPS   "  # fila de categoría
        est = s.detectar_estructura(ws)
        s.sincronizar(ws, est)
        assert ws["A3"].value == "  UPS   "

    def test_codigos_invalidos_generan_advertencia_y_no_se_modifican(self, act):
        ws = self._ws(act, marcas={(0, 0): "X"})
        est = s.detectar_estructura(ws)
        cambios, adv = s.sincronizar(ws, est)
        assert ws["H4"].value == "X"
        assert len(adv) == 1 and adv[0].startswith("H4:")

    def test_datos_nuevos_reemplazan_solo_lo_que_cambia(self, act):
        ws = self._ws(act, marcas={(0, 0): "P"})
        est = s.detectar_estructura(ws)
        nuevos = {(4, "C"): "Luis", (4, "H"): "E", (4, "D"): "MENSUAL"}  # D igual -> sin cambio
        cambios, _ = s.sincronizar(ws, est, nuevos)
        assert {c.celda for c in cambios} == {"C4", "H4"}
        assert ws["C4"].value == "Luis"
        assert ws["H4"].value == "E"

    def test_dato_nuevo_vacio_borra_el_estado(self, act):
        ws = self._ws(act, marcas={(0, 0): "E"})
        est = s.detectar_estructura(ws)
        s.sincronizar(ws, est, {(4, "H"): ""})
        assert ws["H4"].value is None

    def test_segunda_corrida_es_idempotente(self, act):
        ws = self._ws(act, marcas={(0, 0): "e"}, nombre=" A  b ")
        est = s.detectar_estructura(ws)
        s.sincronizar(ws, est)
        cambios, _ = s.sincronizar(ws, est)
        assert cambios == []


class TestCargarDatosXml:
    def test_aplana_y_omite_filas_de_totales(self, tmp_path):
        xml = tmp_path / "d.xml"
        xml.write_text(
            '<Datos><Fila numero="4"><Celda columna="A">Nueva\n</Celda><Celda columna="H">E</Celda></Fila>'
            '<Fila numero="9"><Celda columna="A">Total</Celda></Fila></Datos>',
            encoding="utf-8",
        )
        datos = s.cargar_datos_xml(xml, {9})
        assert datos == {(4, "A"): "Nueva", (4, "H"): "E"}

    def test_celda_vacia(self, tmp_path):
        xml = tmp_path / "d.xml"
        xml.write_text('<Datos><Fila numero="4"><Celda columna="H"/></Fila></Datos>', encoding="utf-8")
        assert s.cargar_datos_xml(xml, set()) == {(4, "H"): ""}


def test_cambio_str():
    assert str(s.Cambio("A1", "x", "y", "motivo")) == "A1: 'x' -> 'y' [motivo]"


def test_constantes_coinciden_con_el_parser():
    from app import plan_trabajo_parser as p

    assert s.WEEK_FIRST == p.WEEK_FIRST and s.WEEK_LAST == p.WEEK_LAST
    assert s.VALID_ESTADOS == set(p.ESTADO_MAP)
